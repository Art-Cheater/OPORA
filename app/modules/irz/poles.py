"""Street-lighting poles: Excel import, list/search and bbox GeoJSON for the IRZ map.

In опоры.xlsx every row is one luminaire, so a pole with several luminaires repeats its
number; such rows are merged into one pole whose quantity is the sum of the rows.
"""

from __future__ import annotations

import io
import math
from collections import Counter
from pathlib import Path
from typing import IO

from app.core.geo.bbox import BboxError, parse_bbox
from app.extensions import db
from app.models.irz import LightPole

SHEET_COLUMNS = {
    "Номер опоры": "pole_number",
    "Название светильника": "luminaire_name",
    "Широта": "latitude",
    "Долгота": "longitude",
    "Кол-во, шт": "quantity",
}
FIELDS = ("luminaire_name", "latitude", "longitude", "quantity")
COORDINATE_TOLERANCE = 1e-5
MAP_LIMIT = 2500
PER_PAGE = 50


def default_poles_path(root: str | Path) -> Path:
    return Path(root) / "опоры.xlsx"


def _header(value: object) -> str:
    return " ".join(str(value or "").split()).casefold().replace("ё", "е")


def normalize_pole_number(value: object) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        if value.is_integer():
            value = int(value)
    text = "".join(str(value).split())
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text[:40] or None


def _number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        value = value.strip().replace(",", ".")
        if not value:
            return None
    number = float(value)
    if not math.isfinite(number):
        raise ValueError("not finite")
    return number


def _coordinates(lat_raw: object, lon_raw: object) -> tuple[float | None, float | None, str | None]:
    try:
        latitude, longitude = _number(lat_raw), _number(lon_raw)
    except (TypeError, ValueError):
        return None, None, "невалидные координаты"
    if latitude is None and longitude is None:
        return None, None, "координаты не указаны"
    if latitude is None or longitude is None:
        return None, None, "указана только одна координата"
    if not (-90 <= latitude <= 90 and -180 <= longitude <= 180) or (latitude == 0 and longitude == 0):
        return None, None, "координаты вне допустимого диапазона"
    return latitude, longitude, None


def _quantity(value: object) -> int | None:
    number = _number(value)
    if number is None:
        return None
    if not number.is_integer() or number < 0:
        raise ValueError("not a count")
    return int(number)


def _luminaires(names: Counter) -> str | None:
    parts = [f"{name} × {count}" if count > 1 else name for name, count in names.items()]
    text = "; ".join(parts)
    return text[:500] or None


def _read_rows(source) -> tuple[list[tuple[int, dict]], int]:
    import openpyxl

    if hasattr(source, "read"):
        source = io.BytesIO(source.read())
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None) or ()
        keys = {_header(name): key for name, key in SHEET_COLUMNS.items()}
        positions = {keys[_header(cell)]: index for index, cell in enumerate(header) if _header(cell) in keys}
        missing = [name for name, key in SHEET_COLUMNS.items() if key not in positions]
        if missing:
            raise ValueError("Не найдены столбцы: " + ", ".join(missing))
        result, total = [], 0
        for number, row in enumerate(rows, start=2):
            if row is None or all(cell in (None, "") for cell in row):
                continue
            total += 1
            result.append((number, {key: row[index] if index < len(row) else None for key, index in positions.items()}))
        return result, total
    finally:
        workbook.close()


def _group(rows: list[tuple[int, dict]], report: dict) -> dict[str, dict]:
    poles: dict[str, dict] = {}
    for number, values in rows:
        pole_number = normalize_pole_number(values.get("pole_number"))
        if pole_number is None:
            report["skipped"] += 1
            report["skip_reasons"]["пустой номер опоры"] += 1
            continue

        def warn(text: str, number=number, pole_number=pole_number) -> None:
            report["warnings"].append(f"строка {number} (опора № {pole_number}): {text}")

        try:
            quantity = _quantity(values.get("quantity"))
        except (TypeError, ValueError):
            report["errors"] += 1
            warn("невалидное количество, строка не учтена в количестве")
            quantity = None
        latitude, longitude, problem = _coordinates(values.get("latitude"), values.get("longitude"))
        name = " ".join(str(values.get("luminaire_name") or "").split())
        pole = poles.get(pole_number)
        if pole is None:
            pole = poles[pole_number] = {"names": Counter(), "quantity": None, "latitude": None, "longitude": None,
                                         "coordinates_problem": None, "row": number}
        else:
            report["merged"] += 1
        if name:
            pole["names"][name] += 1
        if quantity is not None:
            pole["quantity"] = (pole["quantity"] or 0) + quantity
        if latitude is None:
            pole["coordinates_problem"] = pole["coordinates_problem"] or problem
            continue
        if pole["latitude"] is None:
            pole["latitude"], pole["longitude"] = latitude, longitude
        elif abs(pole["latitude"] - latitude) > COORDINATE_TOLERANCE or abs(pole["longitude"] - longitude) > COORDINATE_TOLERANCE:
            warn("координаты отличаются от первой строки этой опоры, оставлены координаты первой строки")
    for pole_number, pole in poles.items():
        if pole["latitude"] is None:
            report["warnings"].append(f"строка {pole['row']} (опора № {pole_number}): "
                                      f"{pole['coordinates_problem'] or 'нет координат'}, опора сохранена без координат и не показывается на карте")
    return poles


def import_poles(source: str | Path | IO[bytes], *, user_id=None, dry_run: bool = False) -> dict:
    """Idempotent upsert by pole_number; poles missing from the file are left untouched."""
    report = {"total": 0, "poles": 0, "inserted": 0, "updated": 0, "unchanged": 0, "skipped": 0, "errors": 0,
              "merged": 0, "skip_reasons": Counter(), "warnings": []}
    rows, report["total"] = _read_rows(source)
    grouped = _group(rows, report)
    report["poles"] = len(grouped)
    existing = {pole.pole_number: pole for pole in db.session.scalars(db.select(LightPole))}
    for pole_number, item in grouped.items():
        data = {"luminaire_name": _luminaires(item["names"]), "latitude": item["latitude"],
                "longitude": item["longitude"], "quantity": item["quantity"]}
        pole = existing.get(pole_number)
        if pole is None:
            db.session.add(LightPole(pole_number=pole_number, created_by=user_id, updated_by=user_id, **data))
            report["inserted"] += 1
            continue
        if pole.deleted_at is None and all(getattr(pole, key) == value for key, value in data.items()):
            report["unchanged"] += 1
            continue
        for key, value in data.items():
            setattr(pole, key, value)
        pole.deleted_at = None
        pole.updated_by = user_id
        report["updated"] += 1
    if dry_run:
        db.session.rollback()
    else:
        db.session.commit()
    report["skip_reasons"] = dict(report["skip_reasons"])
    return report


def serialize_pole(pole: LightPole) -> dict:
    return {
        "id": str(pole.id),
        "pole_number": pole.pole_number,
        "luminaire_name": pole.luminaire_name,
        "quantity": pole.quantity,
        "lat": pole.latitude,
        "lon": pole.longitude,
    }


def get_pole(pole_id) -> LightPole | None:
    return db.session.scalar(db.select(LightPole).where(LightPole.active_filter(), LightPole.id == pole_id))


def search_poles(query: str = "", *, page: int = 1, per_page: int = PER_PAGE) -> tuple[list[LightPole], dict]:
    stmt = db.select(LightPole).where(LightPole.active_filter())
    needle = " ".join(str(query or "").split())
    if needle:
        escaped = needle.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        pattern = f"%{escaped}%"
        stmt = stmt.where(db.or_(LightPole.pole_number.ilike(pattern, escape="\\"),
                                 LightPole.luminaire_name.ilike(pattern, escape="\\")))
    total = db.session.scalar(db.select(db.func.count()).select_from(stmt.subquery())) or 0
    pages = max((total + per_page - 1) // per_page, 1)
    page = min(max(page, 1), pages)
    items = list(db.session.scalars(stmt.order_by(db.func.length(LightPole.pole_number), LightPole.pole_number)
                                    .offset((page - 1) * per_page).limit(per_page)))
    return items, {"page": page, "pages": pages, "per_page": per_page, "total": total}


def parse_map_bbox(args) -> tuple[float, float, float, float]:
    """`bbox=minLon,minLat,maxLon,maxLat` (as MapLibre gives it) or the shared min_lat/max_lat/... keys."""
    raw = args.get("bbox")
    if raw:
        parts = str(raw).split(",")
        if len(parts) != 4:
            raise BboxError("bbox: нужно четыре числа minLon,minLat,maxLon,maxLat.")
        min_lon, min_lat, max_lon, max_lat = parts
        return parse_bbox({"min_lat": min_lat, "max_lat": max_lat, "min_lon": min_lon, "max_lon": max_lon})
    bbox = parse_bbox(args)
    if bbox is None:
        raise BboxError("Для слоя опор нужен bbox видимой области карты.")
    return bbox


def poles_geojson(bbox: tuple[float, float, float, float], *, limit: int = MAP_LIMIT) -> dict:
    min_lat, max_lat, min_lon, max_lon = bbox
    rows = list(db.session.execute(
        db.select(LightPole.id, LightPole.pole_number, LightPole.luminaire_name, LightPole.quantity,
                  LightPole.latitude, LightPole.longitude)
        .where(LightPole.active_filter(), LightPole.latitude.between(min_lat, max_lat),
               LightPole.longitude.between(min_lon, max_lon))
        .order_by(LightPole.pole_number).limit(limit + 1)
    ))
    features = [{
        "type": "Feature",
        "geometry": {"type": "Point", "coordinates": [row.longitude, row.latitude]},
        "properties": {"id": str(row.id), "pole_number": row.pole_number, "luminaire_name": row.luminaire_name,
                       "quantity": row.quantity, "lat": row.latitude, "lon": row.longitude},
    } for row in rows[:limit]]
    return {"type": "FeatureCollection", "features": features, "truncated": len(rows) > limit, "limit": limit}
